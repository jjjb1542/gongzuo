#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify a filled 《房产土地租金收缴专项核查底稿》.

Usage:
    python verify_workingpaper.py <xlsx> [--keywords 潮玩,九街] [--ledger-first 9] [--no-wording]

Checks (exit code 1 if any FAIL):
  1. 5 张表齐备：审计资料清单／审计程序表／欠租统计台账／合同与财务对账差异明细表／疑点清单
  2. 台账数据行：序号非空、无"序号空但其它列有值"的断行
  3. 「应收租金（元）」vs「合同测算累计应收租金」：不一致不判错，但差异明细表必须有对应登记
     （该行「对应台账序号」= 台账序号，且「差异类型」含"台账"）
  4. 日期：合同起止为文本 YYYY-MM-DD；欠租起始月份为 YYYY-MM
  5. 已出租行：减免／审批／催收／催收记录四栏勾选文本非空
  6. 「备注」含"来源"；含来源但无行号 → WARN
  7. 金额列 number_format == '#,##0.00'（WARN）
  8. 资料清单：已填「是否提供」的行，「资料载体类型」须命中四枚举；同一文件不得重复登记
  9. 疑点清单：「核查建议」非空；「事实依据」可溯源（行号或《文件名》）；「疑点」不含定性词
 10. 异常资料：备注含"文件加密损坏"的，疑点清单必须有包含该文件名的记录
 11. 5 张表「核查基准日」一致
 12. 正文表述 V1-V4（全 5 表逐格，含大段说明文本）
 13. 残留关键词（可选，--keywords）
 14. 文件可被 openpyxl 重新打开
"""
import argparse
import os
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

SHEET_MATERIAL = "审计资料清单"
SHEET_PROCEDURE = "审计程序表"
SHEET_LEDGER = "欠租统计台账"
SHEET_RECON = "合同与财务对账差异明细表"
SHEET_ISSUES = "疑点清单"
CORE_SHEETS = [SHEET_MATERIAL, SHEET_PROCEDURE, SHEET_LEDGER, SHEET_RECON, SHEET_ISSUES]

# 台账列 → 字段中文名（仅用于输出可读的报错）
LNAME = {
    1: "序号", 4: "资产名称", 12: "资产状态", 15: "合同开始时间", 16: "合同终止时间",
    17: "应收租金（元）", 24: "合同测算累计应收租金", 28: "欠租起始月份",
    30: "是否存在租金减免", 31: "减免有无审批资料", 32: "是否开展租金催收",
    33: "催收记录是否完整", 34: "备注（合同到期占用/其他情况）",
}
AMOUNT_COLS = [17, 18, 19, 20, 24, 27]  # 应收/已收/未交/欠缴/合同测算/欠租金额
CARRIER_ENUM = ("原件扫描件", "复印件拍照件", "截图", "电子文档")

# ---------------- V1-V4 表述规范 ----------------
DOC_WORDS = ("合同", "评估", "测绘", "招租", "竞价", "决议")
LIST_SEP = set("、，,／/；;。： 　")
BOX_CHARS = ("□", "☑", "■")
# 这两列写的是"资料名称（模板固定）"与"差异类型枚举"，不是文件引用，不卡 V3
V3_EXEMPT_COLS = {(SHEET_MATERIAL, "C"), (SHEET_RECON, "H")}

WORDING_RX = [
    ("V1", "用 Excel 列字母指代字段", re.compile(r"(?<![A-Za-z])[A-Z]{1,2}[ \t]*列")),
    ("V2", "资产简称（N号台账／N号M条）",
     re.compile(r"\d+[ \t]*号[ \t]*台账|\d+[ \t]*号[ \t]*\d+[ \t]*条")),
    ("V3", "形成主体只写机构类别",
     re.compile(r"^(?:评估|咨询|测绘|估价)(?:机构|公司|单位)$")),
    ("V3", "文件裸称（未提供+类别词）",
     re.compile(r"(?:未提供|未见|缺少|缺|无)(?:合同|评估报告|测绘报告|招租方案|竞价结果|董事会决议)"
                r"(?![报案类件价量号方机])")),
    ("V4", "模糊代词（该资产／上述合同…）",
     re.compile(r"该资产|上述(?:合同|资料|文件)|前述(?:合同|资料|文件)")),
    ("V4", "裸写 N 号标的",
     re.compile(r"(?:^|[ \t，。；：、\n])[0-9一二三四五六七八九十]+[ \t]*号[ \t]*标的")),
]

QUAL_WORDS = ("错报", "违规", "虚假", "舞弊", "侵占", "挪用", "贪污", "私分")

results = []


def ok(m):
    results.append(("OK", m))


def warn(m):
    results.append(("WARN", m))


def fail(m):
    results.append(("FAIL", m))


def num(v):
    return isinstance(v, (int, float))


def bare_doc_words(s):
    """未加《》且以分隔符/句尾收束的文件类别词（V3）。"""
    out = []
    for w in DOC_WORDS:
        i = s.find(w)
        while i >= 0:
            nxt = s[i + len(w)] if i + len(w) < len(s) else ""
            if nxt == "" or nxt in LIST_SEP:
                out.append(w)
            i = s.find(w, i + 1)
    return out


def ledger_rows(ws, first):
    """按序号列扫数据行；返回 (序号非空的行号列表, 断行列表)。"""
    rows, gaps = [], []
    for r in range(first, ws.max_row + 1):
        a = ws.cell(r, 1).value
        others = any(ws.cell(r, c).value not in (None, "") for c in range(2, 35))
        if a is None or str(a).strip() == "":
            if others:
                gaps.append(r)
            continue
        rows.append(r)
    return rows, gaps


# ---------------- 各项校验 ----------------

def check_sheets(wb):
    for s in CORE_SHEETS:
        if s in wb.sheetnames:
            ok("sheet 存在：%s" % s)
        else:
            fail("缺少 sheet：%s（模板 4 张 + 脚本新建《疑点清单》，必须合计 5 张）" % s)


def check_ledger(wb, first):
    if SHEET_LEDGER not in wb.sheetnames:
        return [], []
    ws = wb[SHEET_LEDGER]
    rows, gaps = ledger_rows(ws, first)
    for r in gaps:
        fail("台账 R%d：序号为空但其它列有值（数据断行）" % r)
    ok("台账数据行数=%d（R%d 起）" % (len(rows), first))

    for r in rows:
        no = ws.cell(r, 1).value
        q, x = ws.cell(r, 17).value, ws.cell(r, 24).value
        if num(q) and num(x):
            if abs(q - x) < 0.005:
                ok("台账序号%s：应收租金=合同测算=%s" % (no, q))
        elif num(q) or num(x):
            warn("台账序号%s：应收租金=%r、合同测算=%r 有一方非数值，无法比对；"
                 "缺失方须在疑点清单登记" % (no, q, x))
        # 日期
        for c in (15, 16):
            v = ws.cell(r, c).value
            if v not in (None, "", "—") and not isinstance(v, str):
                fail("台账序号%s 的%s非文本：%r（应为文本 YYYY-MM-DD）"
                     % (no, LNAME[c], v))
        ab = ws.cell(r, 28).value
        if ab not in (None, "", "—") and not isinstance(ab, str):
            fail("台账序号%s 的欠租起始月份非文本：%r（应为文本 YYYY-MM）" % (no, ab))
        elif isinstance(ab, str) and ab.strip() not in ("", "—"):
            if not re.match(r"^\d{4}-\d{2}$", ab.strip()):
                fail("台账序号%s 的欠租起始月份格式=%r（应为 YYYY-MM）" % (no, ab))
        # 勾选框
        if str(ws.cell(r, 12).value or "").replace(" ", "").strip() == "已出租":
            for c in (30, 31, 32, 33):
                if not str(ws.cell(r, c).value or "").strip():
                    warn("台账序号%s 的「%s」为空" % (no, LNAME[c]))
        # 来源溯源
        ah = str(ws.cell(r, 34).value or "")
        if not ah:
            fail("台账序号%s 的「%s」为空：须写事实、差异原因与来源" % (no, LNAME[34]))
        elif "来源" not in ah:
            fail("台账序号%s 的「%s」未含“来源：”标注，无法溯源" % (no, LNAME[34]))
        elif not re.search(r"第\s*\d+\s*行", ah):
            warn("台账序号%s 的来源未写台账行号（如「第 61 行」）" % no)
        # 金额格式
        for c in AMOUNT_COLS:
            if num(ws.cell(r, c).value) and ws.cell(r, c).number_format != "#,##0.00":
                warn("台账序号%s 的「%s」格式=%s（期望 #,##0.00）"
                     % (no, LNAME.get(c, get_column_letter(c)), ws.cell(r, c).number_format))
    return rows, gaps


def check_diff_registered(wb, ledger_row_ids):
    """Q≠X 不判错；但差异明细表必须登记对应记录。"""
    if SHEET_LEDGER not in wb.sheetnames or SHEET_RECON not in wb.sheetnames:
        return
    ws = wb[SHEET_LEDGER]
    need = {}
    for r in ledger_row_ids:
        q, x = ws.cell(r, 17).value, ws.cell(r, 24).value
        if num(q) and num(x) and abs(q - x) >= 0.005:
            need[str(ws.cell(r, 1).value).strip()] = (r, q, x)

    registered = set()
    ws2 = wb[SHEET_RECON]
    for r in range(9, ws2.max_row + 1):
        j = ws2.cell(r, 10).value
        h = str(ws2.cell(r, 8).value or "")
        if j in (None, ""):
            continue
        key = str(int(j)) if num(j) else str(j).strip()
        if "台账" in h:
            registered.add(key)

    if not need:
        ok("台账「应收租金」与「合同测算累计应收租金」逐行一致，无需差异登记")
        return
    for key, (r, q, x) in sorted(need.items()):
        if key in registered:
            ok("台账序号%s：应收租金 %s ≠ 合同测算 %s，已在差异明细表登记" % (key, q, x))
        else:
            fail("台账序号%s：应收租金 %s ≠ 合同测算 %s，差异明细表无对应登记"
                 "（须在该表「对应台账序号」填 %s、「差异类型」注明台账与合同测算差异）"
                 % (key, q, x, key))


def check_material_list(wb):
    if SHEET_MATERIAL not in wb.sheetnames:
        return
    ws = wb[SHEET_MATERIAL]
    seen = {}
    for r in range(8, 19):
        d = ws.cell(r, 4).value
        f = ws.cell(r, 6).value
        name = str(ws.cell(r, 3).value or "").strip()
        if d not in (None, ""):
            fs = str(f or "").strip()
            if fs not in CARRIER_ENUM:
                if not (str(d).strip() == "×" and fs in ("—", "-")):
                    fail("资料清单第%d行「资料载体类型」=%r，不在四枚举内"
                         "（原件扫描件／复印件拍照件／截图／电子文档）" % (r, fs))
        if name and d not in (None, ""):
            seen.setdefault(name, []).append(r)
    for name, rs in seen.items():
        if len(rs) > 1 and not re.search(r"补充协议|补充合同|第\s*[2-9]\s*版|修订", name):
            fail("资料清单同一文件登记多次：%s（第 %s 行）；一合同多行只登记 1 次"
                 % (name, "、".join(str(x) for x in rs)))


def check_issues(wb):
    if SHEET_ISSUES not in wb.sheetnames:
        return
    ws = wb[SHEET_ISSUES]
    head = [str(ws.cell(4, c).value or "").strip() for c in range(1, 6)]
    if head[3] != "核查建议":
        fail("疑点清单第 4 行表头应为「序号｜疑点（客观事实）｜事实依据｜核查建议｜资料充分性」，"
             "实际=%s" % head)
    n = 0
    for r in range(5, ws.max_row + 1):
        no = ws.cell(r, 1).value
        if no in (None, ""):
            continue
        n += 1
        fact = str(ws.cell(r, 2).value or "")
        basis = str(ws.cell(r, 3).value or "")
        adv = str(ws.cell(r, 4).value or "").strip()
        if adv in ("", "—", "-", "待补", "待核实", "待人工"):
            fail("疑点清单第%d行（序号%s）「核查建议」为空或占位，须写可执行建议" % (r, no))
        if not (re.search(r"第\s*\d+\s*行", basis) or re.search(r"序号\s*\d+", basis)
                or "《" in basis):
            fail("疑点清单第%d行（序号%s）「事实依据」不可溯源：须含台账行号或《文件名》"
                 % (r, no))
        for w in QUAL_WORDS:
            if w in fact:
                fail("疑点清单第%d行（序号%s）「疑点」出现定性词“%s”，只作客观描述"
                     % (r, no, w))
    ok("疑点清单条数=%d" % n)


def check_abnormal_pdf(wb):
    if SHEET_MATERIAL not in wb.sheetnames or SHEET_ISSUES not in wb.sheetnames:
        return
    ws = wb[SHEET_MATERIAL]
    iss = wb[SHEET_ISSUES]
    all_issue_text = "\n".join(
        str(iss.cell(r, c).value or "") for r in range(5, iss.max_row + 1) for c in range(1, 6))
    for r in range(8, 19):
        e = str(ws.cell(r, 5).value or "")
        if "文件加密损坏" not in e:
            continue
        name = str(ws.cell(r, 3).value or "").strip()
        key = re.sub(r"^\d+[.、]\s*", "", name)
        if key and key not in all_issue_text:
            fail("资料清单第%d行标注“文件加密损坏”，疑点清单未见该文件（%s）的记录" % (r, key))
        else:
            ok("异常资料已留痕：%s" % key)


def check_baseline(wb):
    pat = re.compile(r"核查基准日[：:]\s*([0-9]{4})\s*[年/\-.]\s*([0-9]{1,2})\s*[月/\-.]\s*([0-9]{1,2})")
    found = {}
    for sn in (SHEET_MATERIAL, SHEET_PROCEDURE, SHEET_LEDGER, SHEET_RECON):
        if sn not in wb.sheetnames:
            continue
        m = pat.search(str(wb[sn]["A6"].value or ""))
        if m:
            found[sn] = "%s-%02d-%02d" % (m.group(1), int(m.group(2)), int(m.group(3)))
        else:
            fail("%s 的核查基准日缺失或格式异常" % sn)
    if SHEET_ISSUES in wb.sheetnames:
        m = pat.search(str(wb[SHEET_ISSUES]["A2"].value or ""))
        if m:
            found[SHEET_ISSUES] = "%s-%02d-%02d" % (m.group(1), int(m.group(2)), int(m.group(3)))
        else:
            fail("%s 的核查基准日缺失或格式异常" % SHEET_ISSUES)
    if found:
        vals = set(found.values())
        if len(vals) == 1:
            ok("核查基准日一致：%s（%d 张表）" % (list(vals)[0], len(found)))
        else:
            fail("核查基准日不一致：%s" % found)


def check_wording(wb):
    hits = {"V1": 0, "V2": 0, "V3": 0, "V4": 0}
    for sn in CORE_SHEETS:
        if sn not in wb.sheetnames:
            continue
        for row in wb[sn].iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                s = str(cell.value)
                if any(b in s for b in BOX_CHARS):
                    continue  # 勾选框/枚举串不判
                if s.lstrip().startswith("字段说明"):
                    continue  # 模板自带的字段定义说明区，非本底稿正文
                skip_v3 = (sn, cell.column_letter) in V3_EXEMPT_COLS
                for code, label, rx in WORDING_RX:
                    if code == "V3" and skip_v3:
                        continue
                    m = rx.search(s)
                    if m:
                        hits[code] += 1
                        fail("表述不规范 %s（%s）于 [%s]%s：%s"
                             % (code, label, sn, cell.coordinate, m.group(0)))
                if skip_v3:
                    continue
                for w in bare_doc_words(s):
                    hits["V3"] += 1
                    fail("表述不规范 V3（文件裸称，应写《文件名》+编号+页码）于 [%s]%s：%s"
                         % (sn, cell.coordinate, w))
    if not any(hits.values()):
        ok("表述体检通过（全 5 表逐格；V1 列字母 / V2 资产简称 / V3 文件代称 / V4 模糊代词 均为 0）")
    else:
        ok("表述体检统计：V1=%d V2=%d V3=%d V4=%d（修法见 references/conventions-and-checks.md）"
           % (hits["V1"], hits["V2"], hits["V3"], hits["V4"]))


def check_keywords(wb, keywords):
    hits = 0
    for sn in wb.sheetnames:
        for row in wb[sn].iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                s = str(cell.value)
                for k in keywords:
                    if k and k in s:
                        fail("残留关键词 '%s' 于 [%s]%s" % (k, sn, cell.coordinate))
                        hits += 1
    if hits == 0:
        ok("无残留关键词")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--keywords", default="", help="comma separated residual keywords")
    ap.add_argument("--ledger-first", type=int, default=9)
    ap.add_argument("--no-wording", action="store_true", help="skip the V1-V4 wording check")
    args = ap.parse_args()

    path = os.environ.get("XL", args.xlsx)
    if not os.path.exists(path):
        print("ERROR: not found:", path)
        sys.exit(2)

    wb = openpyxl.load_workbook(path)

    check_sheets(wb)
    rows, _ = check_ledger(wb, args.ledger_first)
    check_diff_registered(wb, rows)
    check_material_list(wb)
    check_issues(wb)
    check_abnormal_pdf(wb)
    check_baseline(wb)

    if args.no_wording:
        warn("表述体检已跳过（--no-wording，不推荐）")
    else:
        check_wording(wb)

    kws = [k.strip() for k in args.keywords.split(",") if k.strip()]
    if kws:
        check_keywords(wb, kws)

    print("=" * 72)
    for lvl, m in results:
        print("[%s] %s" % (lvl, m))
    print("=" * 72)
    nfail = sum(1 for l, _ in results if l == "FAIL")
    print("FAIL=%d  WARN=%d" % (nfail, sum(1 for l, _ in results if l == "WARN")))
    sys.exit(1 if nfail else 0)


if __name__ == "__main__":
    main()
