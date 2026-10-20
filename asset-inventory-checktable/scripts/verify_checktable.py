#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify a 《资产清查检查程序表》 workbook.

Usage:
    python verify_checktable.py <xlsx> [--keywords 旧项目名,错误词] [--no-wording]

Checks (exit code 1 if any FAIL):
  1. five sheets present, in order
  2. header column counts 8 / 10 / 24 / 6 / 2
  3. 检查程序表 conclusion column uses only 符合/部分符合/不符合
  4. 合计 row statistics match the per-procedure conclusions
  5. 资产台账核对: 面积差异 == 测绘参考面积 - 台账面积 (when both numeric)
  6. money/area columns use '#,##0.00'
  7. residual keywords absent (optional)
  8. wording discipline V1-V5: applied to EVERY non-empty cell of all five
     sheets -- including header lines, note rows and long narrative blocks.
     Only 资料清单 C/D (file name / file type) skip V3 and V5.
  9. 表4「涉及资产」列不得带行号／序号／台账表名（应用「坐落＋业态」简称）
 10. file reopens without error
"""
import argparse
import os
import re
import sys
from collections import Counter

import openpyxl

SHEETS = [("检查程序表", 8), ("资料清单", 10), ("资产台账核对", 24),
          ("核查发现与整改建议", 6), ("编制说明", 2)]
MONEY_COLS = (3, 15, 16, 17, 18, 19)
GRADES = ("符合", "部分符合", "不符合")

# V1-V5 覆盖全部 5 张表的所有单元格文本（含表底说明行、编制说明整段、备注与意见栏）。
# 例外只有两类：资料清单的「资料名称」「文件类型」两列不卡 V3/V5——它们写的是文件名/类别名，
# 不是“引用某份文件”，文件名里也可能带 sheet。
V3_EXEMPT = {("资料清单", "C"), ("资料清单", "D")}
V5_EXEMPT = {("资料清单", "C"), ("资料清单", "D")}

# 表4「涉及资产」列（D 列）：只写「坐落＋业态」简称，不得带行号/序号/台账表名。
S4_ASSET_SHEET = "核查发现与整改建议"
S4_ASSET_RX = [
    ("行号", re.compile(r"第[ \t]*\d+[ \t]*行")),
    ("序号", re.compile(r"序号[ \t]*\d+")),
    ("台账表名", re.compile(r"房产表|土地表")),
]

LIST_SEP = set("、，,／/；;。： 　")
DOC_WORDS = ("合同", "评估", "测绘", "招租", "竞价", "决议")

WORDING_RX = [
    ("V1", "列字母指代台账字段", re.compile(r"(?<![A-Za-z])[A-Z]{1,2}[ \t]*列")),
    ("V2", "资产简称（N号台账／N号M条）",
     re.compile(r"\d+[ \t]*号[ \t]*台账|\d+[ \t]*号[ \t]*\d+[ \t]*条")),
    ("V3", "形成主体只写机构类别",
     re.compile(r"^(?:评估|咨询|测绘|估价)(?:机构|公司|单位)$")),
    ("V3", "文件裸称（未提供+类别词）",
     re.compile(r"(?:未提供|未见|缺少|缺|无)(?:合同|评估报告|测绘报告|招租方案|竞价结果|董事会决议)"
                r"(?![报案类件价量号方机])")),
    ("V4", "含糊指代（该资产／上述资料…）",
     re.compile(r"该资产|上述资料|上述文件|前述文件|前述资料|该文件|该合同")),
    ("V4", "含糊指代（N 条记录／中文数字+列）",
     re.compile(r"\d+[ \t]*条[ \t]*记录|[一二三四五六七八九十两][ \t]*列")),
    ("V5", "台账表名写成 sheet", re.compile(r"(?i)sheet")),
]


def bare_doc_words(s):
    """返回未加《》且以分隔符/句尾收束的文件类别词（V3）。"""
    out = []
    for w in DOC_WORDS:
        i = s.find(w)
        while i >= 0:
            nxt = s[i + len(w)] if i + len(w) < len(s) else ""
            if nxt == "" or nxt in LIST_SEP:
                out.append(w)
            i = s.find(w, i + 1)
    return out

res = []


def ok(m):
    res.append(("OK", m))


def warn(m):
    res.append(("WARN", m))


def fail(m):
    res.append(("FAIL", m))


def check_wording(wb):
    """V1-V5 表述体检：遍历 5 张表的全部单元格，命中即 FAIL。"""
    hits = Counter()
    for sn, _ in SHEETS:
        if sn not in wb.sheetnames:
            continue
        for row in wb[sn].iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                s = str(cell.value)
                skip_v3 = (sn, cell.column_letter) in V3_EXEMPT
                skip_v5 = (sn, cell.column_letter) in V5_EXEMPT
                for code, label, rx in WORDING_RX:
                    if code == "V3" and skip_v3:
                        continue
                    if code == "V5" and skip_v5:
                        continue
                    m = rx.search(s)
                    if m:
                        hits[code] += 1
                        fail("表述不规范 " + code + "（" + label + "）于 [" + sn + "]" +
                             cell.coordinate + "：" + m.group(0))
                if skip_v3:
                    continue
                for w in bare_doc_words(s):
                    hits["V3"] += 1
                    fail("表述不规范 V3（文件裸称，应写《文件名》）于 [" + sn + "]" +
                         cell.coordinate + "：" + w)
    if not hits:
        ok("表述体检通过（全表逐格；V1 列字母 / V2 资产简称 / V3 文件代称 / V4 含糊指代 / "
           "V5 台账表名写成 sheet 均为 0）")
    else:
        ok("表述体检统计：V1=" + str(hits["V1"]) + " V2=" + str(hits["V2"]) +
           " V3=" + str(hits["V3"]) + " V4=" + str(hits["V4"]) +
           " V5=" + str(hits["V5"]) + "（修法见 references/conventions.md 第二节）")


def check_s4_asset(wb):
    """表4「涉及资产」列（D）只能用「坐落＋业态」简称，不得带行号／序号／台账表名。"""
    if S4_ASSET_SHEET not in wb.sheetnames:
        return
    ws = wb[S4_ASSET_SHEET]
    bad = []
    for row in ws.iter_rows(min_col=4, max_col=4):
        cell = row[0]
        if cell.value is None:
            continue
        s = str(cell.value)
        for label, rx in S4_ASSET_RX:
            m = rx.search(s)
            if m:
                bad.append(cell.coordinate + "（" + label + "：" + m.group(0) + "）")
    if bad:
        fail("表4「涉及资产」列应写“坐落＋业态”简称（如 圣谕亭巷41号商铺），不得带行号／序号／台账表名：" +
             "；".join(bad))
    else:
        ok("表4「涉及资产」列均为简称（无行号／序号／台账表名）")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--keywords", default="", help="comma separated residual keywords")
    ap.add_argument("--no-wording", action="store_true",
                    help="skip the V1-V5 wording check")
    args = ap.parse_args()

    if not os.path.exists(args.xlsx):
        print("ERROR: not found:", args.xlsx)
        sys.exit(2)
    wb = openpyxl.load_workbook(args.xlsx)

    # 1) sheets
    names = wb.sheetnames
    if names == [s for s, _ in SHEETS]:
        ok("5 张表齐备且顺序正确：" + " / ".join(names))
    else:
        fail("sheet 不符，实际：" + " / ".join(names))

    # 2) header column counts
    for sn, n in SHEETS:
        if sn not in wb.sheetnames:
            fail("缺少 sheet：" + sn)
            continue
        ws = wb[sn]
        hr = 4 if sn == "检查程序表" else 2
        cnt = sum(1 for c in range(1, n + 3) if ws.cell(hr, c).value not in (None, ""))
        if cnt == n:
            ok(sn + " 表头列数=" + str(cnt))
        else:
            fail(sn + " 表头列数=" + str(cnt) + "（期望 " + str(n) + "）")

    # 3/4) procedures + totals
    if "检查程序表" in wb.sheetnames:
        ws = wb["检查程序表"]
        grades, r = [], 5
        while ws.cell(r, 2).value and str(ws.cell(r, 1).value) != "合计":
            g = str(ws.cell(r, 8).value or "").strip()
            grades.append((r, g))
            r += 1
        if not grades:
            warn("检查程序表无程序行")
        else:
            ok("程序行数=" + str(len(grades)))
            badg = [(rr, g) for rr, g in grades if g not in GRADES]
            if badg:
                fail("结论列存在非三档取值：" + str(badg))
            else:
                ok("结论列全部为 符合/部分符合/不符合")
            cg = Counter(g for _, g in grades)
            expect = "符合 " + str(cg["符合"]) + " 项；部分符合 " + str(cg["部分符合"]) + " 项；不符合 " + str(cg["不符合"]) + " 项"
            got = str(ws.cell(r, 6).value or "")
            if expect in got:
                ok("合计统计一致：" + got)
            else:
                fail("合计统计不一致，表内=" + got + "，应为=" + expect)
            if str(ws.cell(r, 1).value) == "合计":
                ok("存在合计行")
            else:
                warn("未找到合计行")

    # 5) area diff arithmetic
    if "资产台账核对" in wb.sheetnames:
        ws = wb["资产台账核对"]
        n, bad, skipped = 0, [], 0
        r = 3
        while ws.cell(r, 1).value and "说明" not in str(ws.cell(r, 1).value):
            ta, sa, diff = ws.cell(r, 3).value, ws.cell(r, 18).value, ws.cell(r, 19).value
            if all(isinstance(x, (int, float)) for x in (ta, sa, diff)):
                if abs((sa - ta) - diff) > 0.005:
                    bad.append((r, ta, sa, diff))
            skip = True
            for x in (ta, sa, diff):
                if isinstance(x, (int, float)):
                    skip = False
            if skip:
                skipped += 1
            n += 1
            r += 1
        ok("资产行数=" + str(n) + ("" if not skipped else "（其中 " + str(skipped) + " 行面积不可比，已跳过）"))
        if bad:
            fail("面积差异算式不符（行, 台账, 测绘, 差异）：" + str(bad))
        else:
            ok("面积差异＝测绘参考面积－台账面积（数值行全部通过）")

    # 6) number formats
    if "资产台账核对" in wb.sheetnames:
        ws = wb["资产台账核对"]
        prob = []
        r = 3
        while ws.cell(r, 1).value and "说明" not in str(ws.cell(r, 1).value):
            for c in MONEY_COLS:
                cell = ws.cell(r, c)
                if isinstance(cell.value, (int, float)) and cell.number_format != "#,##0.00":
                    prob.append(cell.coordinate)
            r += 1
        if prob:
            fail("金额/面积列格式非 #,##0.00：" + ", ".join(prob))
        else:
            ok("金额/面积列均为 #,##0.00")

    # 7) keywords
    kws = [k.strip() for k in args.keywords.split(",") if k.strip()]
    if kws:
        hits = 0
        for sn in wb.sheetnames:
            for row in wb[sn].iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    s = str(cell.value)
                    for k in kws:
                        if k in s:
                            fail("残留关键词 '" + k + "' 于 [" + sn + "]" + cell.coordinate)
                            hits += 1
        if hits == 0:
            ok("无残留关键词")

    # 8) wording discipline
    if args.no_wording:
        warn("表述体检已跳过（--no-wording）")
    else:
        check_wording(wb)
        check_s4_asset(wb)

    print("=" * 72)
    for lv, m in res:
        print("[" + lv + "] " + m)
    print("=" * 72)
    nf = sum(1 for lv, _ in res if lv == "FAIL")
    print("FAIL=" + str(nf) + "  WARN=" + str(sum(1 for lv, _ in res if lv == "WARN")))
    sys.exit(1 if nf else 0)


if __name__ == "__main__":
    main()
