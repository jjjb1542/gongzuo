#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build a 《资产清查检查程序表》 workbook from a JSON spec.

The workbook follows the 5-sheet layout of 新都街道移交资料检查程序表.xlsx:
    检查程序表 / 资料清单 / 资产台账核对 / 核查发现与整改建议 / 编制说明

Usage:
    python checktable.py spec.json
    python checktable.py --template out.xlsx     # blank 5-sheet template
    python checktable.py --example               # print an example spec

Spec schema (rows are positional lists mirroring the sheet columns exactly):
{
  "output": "C:/…/XX资产清查检查程序表.xlsx",
  "title":  "XX资产清查检查程序表",
  "object": "检查对象：…",
  "date":   "编制日期：2026-10-19",
  "basis":  "检查依据：…",
  "scope":  "资料范围：…",
  "procedures": [[序号, 审计程序, 执行内容, 核查要点, 应取证资料,
                  移交资料提供情况, 检查发现, 检查结论(符合/部分符合/不符合)], …],
  "documents": {"title": "…资料清单", "rows": [[10 列], …], "note": "…"},
  "assets":    {"title": "…资产台账核对表", "rows": [[24 列], …], "note": "…"},
  "findings":  {"title": "…", "rows": [[序号, 发现事项, 类型, 涉及资产, 情况说明, 整改建议], …]},
  "notes":     {"title": "编制说明", "rows": [["编制依据", "…"], …]}
}
"""
import argparse
import json
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="B7B7B7")
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="D9E2F3")
SUB_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(size=14, bold=True)
HDR_FONT = Font(bold=True, size=10)
CELL = Font(size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")
CTR = Alignment(horizontal="center", vertical="center")

S1_HEAD = ["序号", "审计程序", "执行内容", "核查要点", "应取证资料",
           "移交资料提供情况", "检查发现", "检查结论"]
S1_W = [5, 15, 34, 34, 13, 46, 50, 9]

S2_HEAD = ["序号", "所属项目", "资料名称", "文件类型", "页数", "形成／出具主体",
           "形成日期", "关键内容摘要", "对应程序", "备注"]
S2_W = [5, 14, 38, 22, 6, 30, 12, 62, 10, 20]

S3_HEAD = ["资产名称", "项目名称", "台账面积（㎡）", "具体位置", "管理公司", "权属公司", "资产类型", "资产业态",
           "取得方式", "取得时间", "权证办理", "资产状态", "资产使用人", "合同起止",
           "应收租金（元）", "已收租金（元）", "未交租金（元）", "测绘参考面积（㎡）", "面积差异（㎡）",
           "差异类型", "差异原因", "取证资料清单", "检查意见", "备注"]
S3_W = [26, 16, 11, 24, 22, 22, 8, 8, 8, 12, 9, 8, 24, 22, 12, 12, 12, 13, 10, 28, 44, 42, 44, 24]

S4_HEAD = ["序号", "发现事项", "类型", "涉及资产", "情况说明", "整改建议"]
S4_W = [5, 24, 8, 24, 54, 52]

S5_HEAD = ["项目", "内容"]
S5_W = [16, 116]

MONEY_COLS = (3, 15, 16, 17, 18, 19)   # 台账面积/租金/测绘面积/面积差异


def _head(ws, row, hdr):
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row, j, h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.border = BD
        c.alignment = WRAP_C


def _box(ws, r0, r1, ncol, ctrcols=()):
    for r in range(r0, r1 + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(r, c)
            cell.font = CELL
            cell.border = BD
            cell.alignment = CTR if c in ctrcols else WRAP


def _widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _merge(ws, r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def build(spec, out):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---------- 表1 检查程序表 ----------
    ws = wb.create_sheet("检查程序表")
    ws["A1"] = spec.get("title", "资产清查检查程序表")
    ws["A1"].font = TITLE_FONT
    ws["A2"] = spec.get("object", "检查对象：")
    ws["E2"] = spec.get("date", "")
    ws["A3"] = spec.get("basis", "检查依据：《资产清查程序表》")
    ws["E3"] = spec.get("scope", "资料范围：")
    for a in ("A1", "A2", "A3", "E2", "E3"):
        ws[a].font = TITLE_FONT if a == "A1" else Font(size=10)
        ws[a].alignment = WRAP
    _merge(ws, 1, 1, 1, 8)
    _merge(ws, 2, 1, 2, 4)
    _merge(ws, 2, 5, 2, 8)
    _merge(ws, 3, 1, 3, 4)
    _merge(ws, 3, 5, 3, 8)
    _head(ws, 4, S1_HEAD)

    procs = spec.get("procedures") or []
    r = 5
    for row in procs:
        for j, v in enumerate(row, 1):
            ws.cell(r, j).value = v
        r += 1
    if procs:
        _box(ws, 5, r - 1, 8, ctrcols=(1, 8))
        for rr in range(5, r):
            ws.row_dimensions[rr].height = 214
        # 合计行：结论统计自动计算
        grades = [str(p[7]).strip() if len(p) > 7 else "" for p in procs]
        n_ok = grades.count("符合")
        n_mid = grades.count("部分符合")
        n_bad = grades.count("不符合")
        ws.cell(r, 1, "合计")
        ws.cell(r, 2, spec.get("total_label", "检查结论统计（共 " + str(len(procs)) + " 项程序）"))
        ws.cell(r, 6, "符合 " + str(n_ok) + " 项；部分符合 " + str(n_mid) + " 项；不符合 " + str(n_bad) + " 项")
        _merge(ws, r, 2, r, 5)
        _merge(ws, r, 6, r, 8)
        _box(ws, r, r, 8, ctrcols=(2, 6))
        for c in range(1, 9):
            ws.cell(r, c).fill = SUB_FILL
            ws.cell(r, c).font = Font(bold=True, size=10)
        ws.row_dimensions[r].height = 26
    _widths(ws, S1_W)
    ws.freeze_panes = "B5"

    # ---------- 表2 资料清单 ----------
    d = spec.get("documents") or {}
    ws2 = wb.create_sheet("资料清单")
    ws2["A1"] = d.get("title", "资料清单")
    ws2["A1"].font = TITLE_FONT
    _merge(ws2, 1, 1, 1, 10)
    _head(ws2, 2, S2_HEAD)
    rows = d.get("rows") or []
    r = 3
    for row in rows:
        for j, v in enumerate(row, 1):
            ws2.cell(r, j).value = v
        r += 1
    if rows:
        _box(ws2, 3, r - 1, 10, ctrcols=(1, 5, 7, 9))
        for rr in range(3, r):
            ws2.row_dimensions[rr].height = 58
    if d.get("note"):
        ws2.cell(r, 1, d["note"])
        _merge(ws2, r, 1, r, 10)
        ws2.cell(r, 1).font = Font(size=9, italic=True)
        ws2.cell(r, 1).alignment = WRAP
        ws2.row_dimensions[r].height = 34
    _widths(ws2, S2_W)
    ws2.freeze_panes = "C3"

    # ---------- 表3 资产台账核对 ----------
    a = spec.get("assets") or {}
    ws3 = wb.create_sheet("资产台账核对")
    ws3["A1"] = a.get("title", "资产台账核对表")
    ws3["A1"].font = TITLE_FONT
    _merge(ws3, 1, 1, 1, 24)
    _head(ws3, 2, S3_HEAD)
    rows = a.get("rows") or []
    r = 3
    for row in rows:
        for j, v in enumerate(row, 1):
            ws3.cell(r, j).value = v
        r += 1
    if rows:
        _box(ws3, 3, r - 1, 24, ctrcols=(3, 7, 8, 9, 10, 11, 12, 18, 19))
        for rr in range(3, r):
            ws3.row_dimensions[rr].height = 200
            for c in MONEY_COLS:
                cell = ws3.cell(rr, c)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
    if a.get("note"):
        ws3.cell(r, 1, a["note"])
        _merge(ws3, r, 1, r, 24)
        ws3.cell(r, 1).font = Font(size=9, italic=True)
        ws3.cell(r, 1).alignment = WRAP
        ws3.row_dimensions[r].height = 56
    _widths(ws3, S3_W)
    ws3.freeze_panes = "C3"

    # ---------- 表4 核查发现与整改建议 ----------
    f = spec.get("findings") or {}
    ws4 = wb.create_sheet("核查发现与整改建议")
    ws4["A1"] = f.get("title", "核查发现与整改建议")
    ws4["A1"].font = TITLE_FONT
    _merge(ws4, 1, 1, 1, 6)
    _head(ws4, 2, S4_HEAD)
    rows = f.get("rows") or []
    r = 3
    for row in rows:
        for j, v in enumerate(row, 1):
            ws4.cell(r, j).value = v
        r += 1
    if rows:
        _box(ws4, 3, r - 1, 6, ctrcols=(1, 3))
        for rr in range(3, r):
            ws4.row_dimensions[rr].height = 62
    _widths(ws4, S4_W)
    ws4.freeze_panes = "A3"

    # ---------- 表5 编制说明 ----------
    n = spec.get("notes") or {}
    ws5 = wb.create_sheet("编制说明")
    ws5["A1"] = n.get("title", "编制说明")
    ws5["A1"].font = TITLE_FONT
    _merge(ws5, 1, 1, 1, 2)
    _head(ws5, 2, S5_HEAD)
    r = 3
    for row in (n.get("rows") or []):
        ws5.cell(r, 1, row[0]).font = Font(bold=True, size=10)
        ws5.cell(r, 1).fill = SUB_FILL
        ws5.cell(r, 1).alignment = WRAP
        ws5.cell(r, 1).border = BD
        ws5.cell(r, 2, row[1] if len(row) > 1 else "").font = CELL
        ws5.cell(r, 2).alignment = WRAP
        ws5.cell(r, 2).border = BD
        r += 1
    _widths(ws5, S5_W)

    wb.save(out)
    return out


def make_template(out):
    """Write a blank 5-sheet template (headers + merges + widths only)."""
    spec = {
        "title": "XX资产清查检查程序表",
        "object": "检查对象：XX（台账「房产」表第X行／序号X）；权属公司：XX",
        "date": "编制日期：",
        "basis": "检查依据：《资产清查程序表》（共9项程序）；《…8.12.xlsx》资产台账",
        "scope": "资料范围：XX资料目录N份（…）",
        "procedures": [],
        "documents": {"title": "XX资料清单", "rows": [], "note": ""},
        "assets": {"title": "XX资产台账核对表", "rows": [], "note": ""},
        "findings": {"title": "核查发现与整改建议（XX）", "rows": []},
        "notes": {"title": "编制说明", "rows": []},
    }
    return build(spec, out)


EXAMPLE = {
    "output": "C:/…/XX资产清查检查程序表.xlsx",
    "title": "XX资产清查检查程序表",
    "object": "检查对象：XX商铺（台账「房产」表第56行／序号151）；权属公司：XX公司",
    "date": "编制日期：2026-10-19",
    "basis": "检查依据：《资产清查程序表》（共9项程序）；《成都香城产业发展集团有限公司8.12_副本.xlsx》资产台账",
    "scope": "资料范围：XX资料目录5份（土地证、《房屋租赁合同》、测绘报告、评估报告、会议纪要）",
    "procedures": [
        [1, "获取并整理资产台账", "从资产管理系统导出房屋、土地资产台账…",
         "1.台账是否存在重复登记、漏登记资产；2.资产编号是否唯一…",
         "资产系统导出资产台账", "已提供…（执行情况）", "①…②…（检查发现）", "部分符合"],
        [2, "收集权属证明文件", "收集对应资产不动产权证书…", "…", "不动产权证原件扫描件",
         "未提供…", "…", "不符合"],
    ],
    "documents": {
        "title": "XX资料清单",
        "rows": [
            [1, "XX项目", "房屋租赁合同.pdf", "房屋租赁合同（扫描件）", 24,
             "出租方：XX／承租方：XX", "2026-04-23", "…主要条款…", "程序2/6", "扫描件，无文本层"],
        ],
        "note": "说明：上表N份为…目录下全部资料；其中扫描件M份。",
    },
    "assets": {
        "title": "XX资产台账核对表",
        "rows": [
            ["XX商铺", "XX项目", 132.79, "XX路41号", "XX公司", "XX公司", "房产", "商铺",
             "无偿划转", "2024-03-11", "无产权证", "已出租", "XX公司", "2026-05-01 至 2027-04-30",
             65333, 32666.5, 32666.5, 132.79, 0,
             "①权证缺失；②账实资料缺失", "①…②…", "【已提供】…【缺失】…", "①…综合：…", "测绘房名2"],
        ],
        "note": "说明：①台账字段取自…；②“面积差异”＝测绘参考面积－台账面积。",
    },
    "findings": {
        "title": "核查发现与整改建议（XX）",
        "rows": [[1, "权证类材料缺失", "缺项", "XX", "…情况…", "…建议…"]],
    },
    "notes": {
        "title": "编制说明",
        "rows": [["编制依据", "…"], ["检查对象", "…"], ["对应资产", "…"], ["检查方法", "…"],
                 ["结论口径", "符合＝资料齐备且无异常；部分符合＝资料齐备但存在瑕疵或差异待处理；不符合＝关键资料缺失或存在重大差异。"],
                 ["主要结论", "…"], ["数据来源", "…"], ["声明", "…"]],
    },
}


def main():
    ap = argparse.ArgumentParser(description="Build 资产清查检查程序表 from a JSON spec.")
    ap.add_argument("spec", nargs="?", help="path to spec.json")
    ap.add_argument("--template", metavar="XLSX", help="write a blank template to XLSX")
    ap.add_argument("--example", action="store_true", help="print an example spec and exit")
    args = ap.parse_args()
    if args.example or not (args.spec or args.template):
        print(json.dumps(EXAMPLE, ensure_ascii=False, indent=2))
        return
    if args.template:
        p = make_template(args.template)
        print("TEMPLATE:", p, os.path.exists(p), os.path.getsize(p))
        return
    with open(args.spec, "r", encoding="utf-8") as fh:
        spec = json.load(fh)
    out = spec["output"]
    build(spec, out)
    print("SAVED:", out, os.path.exists(out), os.path.getsize(out))


if __name__ == "__main__":
    main()
